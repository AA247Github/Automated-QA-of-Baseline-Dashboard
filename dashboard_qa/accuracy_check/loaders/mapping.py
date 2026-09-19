"""Load BI SQL-column mapping workbooks (ID / Column Name / SQL Column Name)."""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook


def _norm(value: object) -> str:
    # Makes workbook headings consistent before they are compared.
    return re.sub(r"\s+", " ", str(value or "").replace("\n", " ")).strip()


CSV_HEADER_ALIASES = {
    "column name",
    "csv column",
    "csv column name",
    "dcs column",
    "capture column",
    "sheet column",
}
# Accepted names for the database-field column in different mapping templates.
SQL_HEADER_ALIASES = {
    "sql column name",
    "sql column",
    "sql field",
    "sql field name",
    "cliics column",
    "database column",
}


def load_mapping(path: Path) -> dict[str, str]:
    """Builds the SQL field to CSV heading lookup used by the calculation engine."""
    wb = load_workbook(path, data_only=True, read_only=True)
    sql_to_csv: dict[str, str] = {}
    try:
        # Uses the first worksheet containing a recognisable pair of columns.
        for sheet in wb.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            header = [_norm(c).lower() for c in rows[0]]
            csv_idx = next((i for i, h in enumerate(header) if h in CSV_HEADER_ALIASES), None)
            sql_idx = next((i for i, h in enumerate(header) if h in SQL_HEADER_ALIASES), None)
            if csv_idx is None:
                csv_idx = next((i for i, h in enumerate(header) if "column name" in h or h == "column"), 1 if len(header) > 1 else None)
            if sql_idx is None:
                sql_idx = next((i for i, h in enumerate(header) if "sql" in h), 2 if len(header) > 2 else None)
            if csv_idx is None or sql_idx is None:
                continue
            for row in rows[1:]:
                if not row or max(csv_idx, sql_idx) >= len(row):
                    continue
                csv_name = _norm(row[csv_idx])
                sql_name = _norm(row[sql_idx])
                if csv_name and sql_name:
                    sql_to_csv[sql_name] = csv_name
            if sql_to_csv:
                break
    finally:
        wb.close()
    if not sql_to_csv:
        raise ValueError(f"No SQL/CSV column pairs found in {path.name}")
    return sql_to_csv


def match_csv_header(csv_name: str, headers: list[str]) -> str | None:
    # Tries an exact heading first, then a looser contained-name match.
    target = _norm(csv_name).lower()
    for h in headers:
        if _norm(h).lower() == target:
            return h
    for h in headers:
        if target and target in _norm(h).lower():
            return h
    return None
